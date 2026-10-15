"""Build submission documents from the latest results.

    python scripts/make_docs.py

Writes docs/architecture.{pdf,png}, docs/GlucoTwin_presentation.{pptx,pdf} and refreshes the
results table inside README.md (between the RESULTS markers).
"""
from __future__ import annotations

import json
import re
import subprocess
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import matplotlib  # noqa: E402

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from matplotlib.patches import FancyArrowPatch, FancyBboxPatch  # noqa: E402

from glucotwin import config as C  # noqa: E402

DOCS = ROOT / "docs"
FIGS = C.REPORTS_DIR / "figures"
BLUE, ORANGE, AQUA, VIOLET, GREEN = "#2a78d6", "#eb6834", "#1baf7a", "#4a3aa7", "#008300"
INK, INK2, MUTED, SURFACE = "#0b0b0b", "#52514e", "#898781", "#fcfcfb"
TEAM = "Team Twin"
COLLEGE = "KIIT, Bhubaneswar"


def metrics() -> dict:
    return json.loads((C.REPORTS_DIR / "metrics.json").read_text())


def hypo_metrics() -> dict | None:
    path = C.REPORTS_DIR / "hypo_metrics.json"
    return json.loads(path.read_text()) if path.exists() else None


def glance_markdown(m: dict) -> str:
    """Headline numbers for the top of the README."""
    ab, ci, fc = m["ablations"], m["auroc_ci"]["CGM only"], m["forecast"]
    ft = ab["GlucoTwin: fusion + digital twin"]
    rd, hy = real_metrics(), hypo_metrics()
    rows = ["| | Synthetic cohort (24 unseen patients) | Real patients (ShanghaiT2DM) |", "|---|---|---|"]
    if rd:
        ra = rd["auroc"]["GlucoTwin (+ digital twin)"]
        rci = rd["diff_ci"]["GlucoTwin (+ digital twin) vs CGM only"]
        rf = rd["forecast"]
        rows += [
            f"| **Spike prediction (AUROC)** | **{ft['auroc']:.3f}** | **{ra['auroc']:.3f}** ({rd['cohort']['subjects']} people) |",
            f"| Gain over CGM alone (95% CI) | {ci['diff']:+.3f} ({ci['lo']:+.3f} to {ci['hi']:+.3f}) | "
            f"{rci['diff']:+.3f} ({rci['lo']:+.3f} to {rci['hi']:+.3f}) |",
            f"| Spikes warned in advance · median warning | {ft['event_sensitivity']:.0%} · {ft['median_lead_min']:.0f} min | "
            f"{ra['event_sensitivity']:.0%} · {ra['median_lead_min']:.0f} min (untuned threshold) |",
            f"| 30-min forecast error: ours vs 'no change' | {fc['30']['GlucoTwin hybrid']['rmse']:.1f} vs "
            f"{fc['30']['Persistence (no change)']['rmse']:.1f} mg/dL | {rf['30']['GlucoTwin hybrid (ML, out-of-fold)']:.1f} vs "
            f"{rf['30']['Persistence (no change)']:.1f} mg/dL |",
            f"| Model trained only on synthetic data, applied to real patients | | AUROC {rd['zero_shot']['auroc']:.3f} |",
        ]
    if hy:
        ha = hy["auroc"]["GlucoTwin (+ digital twin)"]
        rule = hy["auroc"]["Rule: 'urgent low soon' (30-min extrapolation)"]
        mt = hy.get("tradeoff", {}).get("matched", {})
        rows.append(f"| **Low-glucose (<70) prediction (AUROC)** | (too rare in simulation) | **{ha['auroc']:.3f}** "
                    f"(classic rule {rule['auroc']:.3f}) |")
        for b in mt.get("by_alarm_budget", []):
            if b["false_alerts_per_day_budget"] in (0.25, 1.0):
                rows.append(f"| Lows warned with ≤ {b['false_alerts_per_day_budget']:.2f} false alarms/day | | "
                            f"GlucoTwin {b['glucotwin_sensitivity']:.0%} vs best rule setting {b['rule_sensitivity']:.0%} |")
    return "\n".join(rows)


def real_metrics() -> dict | None:
    path = C.REPORTS_DIR / "real_data_metrics.json"
    return json.loads(path.read_text()) if path.exists() else None


# --------------------------------------------------------------------------------------
# Architecture diagram
# --------------------------------------------------------------------------------------
def _box(ax, x, y, w, h, title, body, color):
    ax.add_patch(FancyBboxPatch((x, y), w, h, boxstyle="round,pad=0.008,rounding_size=0.012",
                                fc="white", ec=color, lw=1.6))
    ax.add_patch(FancyBboxPatch((x, y + h - 0.045), w, 0.045, boxstyle="round,pad=0.008,rounding_size=0.012",
                                fc=color, ec=color, lw=1.6))
    ax.text(x + 0.012, y + h - 0.0225, title, color="white", fontsize=10.5, fontweight="bold", va="center")
    ax.text(x + 0.012, y + h - 0.062, body, color=INK2, fontsize=8.4, va="top", linespacing=1.45)


def _arrow(ax, x0, y0, x1, y1, color=MUTED):
    ax.add_patch(FancyArrowPatch((x0, y0), (x1, y1), arrowstyle="-|>", mutation_scale=14, color=color, lw=1.4))


def architecture() -> None:
    fig = plt.figure(figsize=(16, 9))
    ax = fig.add_axes([0, 0, 1, 1])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis("off")
    fig.patch.set_facecolor(SURFACE)
    ax.text(0.02, 0.965, "GlucoTwin: system architecture", fontsize=20, fontweight="bold", color=INK, va="top")
    ax.text(0.02, 0.925, "Hybrid digital twin (mechanistic physiology + machine learning) that predicts hyperglycaemic "
            "excursions in Type 2 Diabetes up to 2 hours ahead", fontsize=11, color=INK2, va="top")

    _box(ax, 0.02, 0.56, 0.19, 0.30, "STATIC  ·  EHR",
         "Demographics, BMI, waist\nDiagnoses & comorbidities\nLabs: HbA1c, FPG, lipids, eGFR\n"
         "Medications (metformin, SU,\n  DPP-4i, SGLT2i, insulin)\nGenetic markers: TCF7L2,\n  KCNQ1, polygenic risk\n"
         "Sources: synthetic generator,\n  Synthea importer", BLUE)
    _box(ax, 0.02, 0.17, 0.19, 0.35, "DYNAMIC  ·  WEARABLES",
         "CGM glucose (5-min)\nSteps / activity\nHeart rate\nHRV (RMSSD)\nSleep sessions & stages\nMeal logs (partial)\n\n"
         "Apple Health / Google Fit /\n  medical-IoT style streams\nSensor lag, noise, dropouts", ORANGE)
    _box(ax, 0.25, 0.42, 0.15, 0.30, "INGESTION",
         "FastAPI REST endpoint\nPOST /observations\nSchema validation\n  (pydantic)\nPer-patient stream\n  buffer (6 days)\n"
         "Same feature code\n  online & offline\nApple Health import\nFHIR R4 export (ABDM)", INK2)
    _box(ax, 0.44, 0.62, 0.25, 0.24, "① MECHANISTIC TWIN",
         "Personalised glucose–insulin ODE\n(Bergman minimal model + gut absorption\n+ exercise + dawn + medication effects)\n"
         "EHR → priors · 7-day MAP calibration\nCGM assimilated every 5 min\nPhysiology forecast & what-if rollouts", VIOLET)
    _box(ax, 0.44, 0.36, 0.25, 0.22, "② CAUSAL FEATURE FUSION",
         "~100 features in 10 groups: glucose trend,\npersonal routine, meals, activity, HR/HRV,\n"
         "sleep, time of day, twin state & residuals,\nclinical EHR, genetics\nNo look-ahead (tested)", GREEN)
    _box(ax, 0.44, 0.10, 0.25, 0.22, "③ ML + EXPLAINABILITY",
         "Gradient-boosted spike classifier\n  (isotonic-calibrated probability)\nGlucose forecasters: +30/60/120 min\n"
         "Group-occlusion attributions →\n  plain-language drivers", AQUA)
    _box(ax, 0.73, 0.42, 0.11, 0.44, "OUTPUTS",
         "P(spike in 2 h)\n+ alert\n\nGlucose forecast\n+ 80% band\n\nTop drivers\n(why)\n\nTwin-tested\nsuggestion\n\n"
         "Twin phenotype\n(SI, β-cell)", ORANGE)
    _box(ax, 0.87, 0.42, 0.115, 0.44, "CLINICIAN UI",
         "Streamlit console\n\nWard overview\n(risk-sorted)\n\nVirtual patient\ntimeline\n\nWhat-if\nsimulator\n\nModel\nperformance", BLUE)

    for y in (0.71, 0.345):
        _arrow(ax, 0.215, y, 0.248, 0.57)
    _arrow(ax, 0.40, 0.57, 0.438, 0.74)
    _arrow(ax, 0.40, 0.57, 0.438, 0.47)
    _arrow(ax, 0.565, 0.62, 0.565, 0.585)
    _arrow(ax, 0.565, 0.36, 0.565, 0.325)
    _arrow(ax, 0.69, 0.74, 0.728, 0.70)
    _arrow(ax, 0.69, 0.21, 0.728, 0.50)
    _arrow(ax, 0.84, 0.64, 0.868, 0.64)

    ax.add_patch(FancyBboxPatch((0.73, 0.06), 0.255, 0.30, boxstyle="round,pad=0.008,rounding_size=0.012",
                                fc="#f3f2ee", ec="#c3c2b7", lw=1))
    ax.text(0.742, 0.345, "OFFLINE PIPELINE  (scripts/build_all.py)", fontsize=9.5, fontweight="bold", color=INK, va="top")
    steps = ["1  Simulate cohort (120 pts × 28 d)", "2  Calibrate twins (days 1–7)", "3  Causal features",
             "4  Train · patient-level split 60/20/20", "5  Evaluate on unseen patients:", "    ablations, alerts, lead time,",
             "    cold start, calibration", "6  Reports, figures, model bundle"]
    for i, s_ in enumerate(steps):
        ax.text(0.745, 0.305 - i * 0.03, s_, fontsize=8.6, color=INK2, va="top")
    ax.text(0.02, 0.03, "All data synthetic (DPDP Act / HIPAA sandbox rules). Research prototype; not a medical device.",
            fontsize=8.5, color=MUTED)
    DOCS.mkdir(exist_ok=True)
    fig.savefig(DOCS / "architecture.pdf")
    fig.savefig(DOCS / "architecture.png", dpi=150)
    plt.close(fig)


# --------------------------------------------------------------------------------------
# README results table
# --------------------------------------------------------------------------------------
def results_markdown(m: dict) -> str:
    ab = m["ablations"]
    order = sorted(ab, key=lambda k: -ab[k]["auroc"])
    lines = ["| Model (test = 24 unseen patients) | AUROC | AUPRC | Spikes caught ≤2 h ahead | Median warning | "
             "Alert precision | False alerts / day |",
             "|---|---|---|---|---|---|---|"]
    for k in order:
        v = ab[k]
        b = "**" if k.startswith("GlucoTwin") else ""
        lines.append(f"| {b}{k}{b} | {v['auroc']:.3f} | {v['auprc']:.3f} | {v['event_sensitivity']:.0%} | "
                     f"{v['median_lead_min']:.0f} min | {v['alert_precision']:.0%} | {v['false_alerts_per_day']:.2f} |")
    fc = m["forecast"]
    lines += ["", "| Glucose forecast RMSE (mg/dL) | +30 min | +60 min | +120 min |", "|---|---|---|---|"]
    for name in ("GlucoTwin hybrid", "Hybrid without twin features", "Mechanistic twin", "Persistence (no change)"):
        lines.append(f"| {name} | " + " | ".join(f"{fc[h][name]['rmse']:.1f}" for h in ("30", "60", "120")) + " |")
    ci = m["auroc_ci"]
    lines += ["", "| AUROC difference: GlucoTwin minus … | Difference | 95% CI (patient bootstrap, 200×) |",
              "|---|---|---|"]
    for name, v in ci.items():
        lines.append(f"| {name} | {v['diff']:+.3f} | {v['lo']:+.3f} to {v['hi']:+.3f} |")
    rd = real_metrics()
    if rd:
        co = rd["cohort"]
        lines += ["", f"**Real patients: ShanghaiT2DM** ({co['subjects']} people with T2D, {co['recordings']} recordings, "
                  f"5-fold cross-validation grouped by person, spike rate {co['spike_rate']:.1%}):", "",
                  "| Model (real data) | AUROC | AUPRC | Spikes caught | Median warning | False alerts / day |",
                  "|---|---|---|---|---|---|"]
        for name, v in rd["auroc"].items():
            lines.append(f"| {name} | {v['auroc']:.3f} | {v['auprc']:.3f} | {v['event_sensitivity']:.0%} | "
                         f"{v['median_lead_min']:.0f} min | {v['false_alerts_per_day']:.2f} |")
        lines.append(f"| Zero-shot: model trained only on synthetic patients | {rd['zero_shot']['auroc']:.3f} | "
                     f"{rd['zero_shot']['auprc']:.3f} | | | |")
        lines += ["", "| Real-data AUROC difference | Difference | 95% CI (bootstrap over people) |", "|---|---|---|"]
        for name, v in rd["diff_ci"].items():
            lines.append(f"| {name} | {v['diff']:+.3f} | {v['lo']:+.3f} to {v['hi']:+.3f} |")
        lines += ["", "| Real-data forecast RMSE (mg/dL) | +30 min | +60 min |", "|---|---|---|"]
        for name in rd["forecast"]["30"]:
            lines.append(f"| {name} | {rd['forecast']['30'][name]:.1f} | {rd['forecast']['60'][name]:.1f} |")
        lines += ["", "_Real-data alerts use a fixed 0.5 probability threshold (no per-cohort tuning)._"]
    hy = hypo_metrics()
    if hy:
        co = hy["cohort"]
        lines += ["", f"**Hypoglycaemia (<70 mg/dL within 2 h), real patients** ({co['people']} people, "
                  f"{co['low_episodes']} low episodes; thresholds chosen by inner cross-validation):", "",
                  "| Model | AUROC | AUPRC | Lows caught | Median warning | False alerts / day |", "|---|---|---|---|---|---|"]
        for name, v in hy["auroc"].items():
            ev = "event_sensitivity" in v
            lines.append(f"| {name} | {v['auroc']:.3f} | {v['auprc']:.3f} | "
                         + (f"{v['event_sensitivity']:.0%} | {v['median_lead_min']:.0f} min | {v['false_alerts_per_day']:.2f} |"
                            if ev else "| | |"))
        lines += ["", "| Hypo AUROC difference | Difference | 95% CI |", "|---|---|---|"]
        for name, v in hy["diff_ci"].items():
            lines.append(f"| {name} | {v['diff']:+.3f} | {v['lo']:+.3f} to {v['hi']:+.3f} |")
    bh = m["by_history"]
    lines += ["", "| AUROC / false alerts per day, by monitoring history | " + " | ".join(bh) + " |",
              "|---|" + "---|" * len(bh)]
    for name in ("CGM only", "Wearables only (dynamic)", "Fusion: EHR + wearables", "GlucoTwin: fusion + digital twin"):
        lines.append(f"| {name} | " + " | ".join(
            f"{bh[b][name]['auroc']:.3f} / {bh[b][name]['false_alerts_per_day']:.2f}" for b in bh) + " |")
    return "\n".join(lines)


def twin_recovery() -> dict:
    """Correlation (log scale) between hidden true parameters and calibrated / EHR-prior estimates."""
    import numpy as np
    import pandas as pd

    from glucotwin import physiology as ph
    from glucotwin import twin

    fit = pd.read_csv(C.MODELS_DIR / "twin_params.csv").set_index("patient_id")
    tru = pd.read_csv(C.RAW_DIR / "true_params.csv").set_index("patient_id").loc[fit.index]
    ehr = pd.read_csv(C.RAW_DIR / "ehr.csv").set_index("patient_id", drop=False).loc[fit.index]
    pri = np.array([twin.prior_params(r) for _, r in ehr.iterrows()])
    out = {}
    for k, i in (("SI", ph.P_SI), ("beta", ph.P_BETA), ("Gb", ph.P_GB)):
        t = np.log(tru[f"true_{k}"])
        out[k] = (float(np.corrcoef(t, np.log(pri[:, i]))[0, 1]), float(np.corrcoef(t, np.log(fit[k]))[0, 1]))
    return out


def takeaways_markdown(m: dict) -> str:
    ab, ci, fc = m["ablations"], m["auroc_ci"], m["forecast"]
    cg, wr, ft = ab["CGM only"], ab["Wearables only (dynamic)"], ab["GlucoTwin: fusion + digital twin"]
    rec = twin_recovery()
    c = ci["CGM only"]
    others = [ci["Wearables only (dynamic)"], ci["Fusion: EHR + wearables"]]
    bound = max(max(abs(v["lo"]), abs(v["hi"])) for v in others)
    h = {k: fc[k]["GlucoTwin hybrid"]["rmse"] for k in ("30", "60", "120")}
    return "\n".join([
        "**Key takeaways (stated with their uncertainty)**",
        f"- **Fusing the dynamic streams is what drives accuracy.** Adding meals, activity, sleep, HR/HRV and personal "
        f"routine to CGM raises AUROC from {cg['auroc']:.3f} to {wr['auroc']:.3f}. GlucoTwin vs CGM-only: "
        f"{c['diff']:+.3f} (95% CI {c['lo']:+.3f} to {c['hi']:+.3f}). False alarms fall from "
        f"{cg['false_alerts_per_day']:.2f} to {ft['false_alerts_per_day']:.2f} per patient-day.",
        f"- **GlucoTwin warns about {ft['event_sensitivity']:.0%} of excursions, with a median "
        f"{ft['median_lead_min']:.0f} minutes' notice.** {ft['alert_precision']:.0%} of its alerts are followed by a real "
        "excursion.",
        f"- **Static EHR and twin features add no measurable discrimination once wearable history exists.** "
        f"All differences fall within ±{bound:.3f} AUROC. A data-efficiency experiment "
        "(`scripts/experiment_data_efficiency.py`, 6–72 training patients) found no gain from the twin either, in this "
        "synthetic setting.",
        "- **What the twin does provide is what a purely statistical model cannot:**",
        "  - *counterfactual* what-if simulation (meal, carbohydrate quality, post-meal walk, sleep) from the patient's "
        "current physiological state;",
        "  - an interpretable phenotype;",
        "  - twin-tested suggestions, such as how much a walk now would lower the peak.",
        "- **The twin recovers hidden physiology** from one week of data, measured against the simulator's ground truth. "
        "Correlation with the true value, EHR priors → calibrated twin:",
        f"  - insulin sensitivity: r = {rec['SI'][0]:.2f} → **{rec['SI'][1]:.2f}**",
        f"  - β-cell response: r = {rec['beta'][0]:.2f} → **{rec['beta'][1]:.2f}**",
        f"  - basal glucose: r = {rec['Gb'][0]:.2f} → **{rec['Gb'][1]:.2f}**",
        f"- **The ML forecaster roughly halves glucose-forecast error** versus the physiology-only twin or a 'no change' "
        f"guess ({h['30']:.0f} / {h['60']:.0f} / {h['120']:.0f} mg/dL RMSE at +30 / 60 / 120 min).",
        *real_takeaways(),
        *hypo_takeaways(),
    ])


def hypo_takeaways() -> list[str]:
    hy = hypo_metrics()
    if not hy:
        return []
    a = hy["auroc"]
    g, rule = a["GlucoTwin (+ digital twin)"], a["Rule: 'urgent low soon' (30-min extrapolation)"]
    d = hy["diff_ci"]["+ doses vs routine & meals"]
    c = hy["diff_ci"]["GlucoTwin vs CGM only"]
    mt = hy.get("tradeoff", {}).get("matched", {})
    match = ""
    bud = {b["false_alerts_per_day_budget"]: b for b in mt.get("by_alarm_budget", [])}
    if 0.25 in bud and 1.0 in bud:
        q, l_ = bud[0.25], bud[1.0]
        match = (f" The alarm trade-off is mixed. When alarms must be rare (≤0.25/day), GlucoTwin warns of "
                 f"{q['glucotwin_sensitivity']:.0%} of lows vs {q['rule_sensitivity']:.0%} for the best rule setting. At "
                 f"≤1 alarm/day the simple rule wins ({l_['rule_sensitivity']:.0%} vs {l_['glucotwin_sensitivity']:.0%}). "
                 "Our design therefore keeps the rule as a safety net for imminent lows and uses GlucoTwin for quieter, "
                 "earlier heads-ups.")
    return [f"- **Lows too, on real patients.** A hypoglycaemia model trained on {hy['cohort']['people']} real people "
            f"({hy['cohort']['low_episodes']} lows) reaches AUROC {g['auroc']:.3f}, versus {rule['auroc']:.3f} for the "
            f"classic 'urgent low soon' rule and {c['diff']:+.3f} over CGM alone (CI {c['lo']:+.3f} to {c['hi']:+.3f})."
            f"{match} Insulin/oral-agent dose data changed AUROC by {d['diff']:+.3f} "
            f"(CI {d['lo']:+.3f} to {d['hi']:+.3f}), so the deployable model works without them."]


def real_takeaways() -> list[str]:
    rd = real_metrics()
    if not rd:
        return []
    a, d, fc, co = rd["auroc"], rd["diff_ci"], rd["forecast"], rd["cohort"]
    best = max(a, key=lambda k: a[k]["auroc"])
    g = d["GlucoTwin (+ digital twin) vs CGM only"]
    e = d["+ EHR (clinical record) vs + meals & insulin doses"]
    t = d["GlucoTwin (+ digital twin) vs + EHR (clinical record)"]
    return [
        f"- **Validated on {co['subjects']} real people with T2D (ShanghaiT2DM).** The fusion gain over CGM alone "
        f"replicates: {g['diff']:+.3f} AUROC (95% CI {g['lo']:+.3f} to {g['hi']:+.3f}). The best real-data configuration "
        f"is *{best}* (AUROC {a[best]['auroc']:.3f}). The model trained only on synthetic patients reaches "
        f"{rd['zero_shot']['auroc']:.3f} without seeing any real data.",
        f"- **On real data the static EHR slightly hurt** ({e['diff']:+.3f}, CI {e['lo']:+.3f} to {e['hi']:+.3f}), "
        f"likely over-fitting with ~80 training people. Twin features were neutral ({t['diff']:+.3f}). The "
        f"**physiology-only twin forecast was worse than 'no change'** "
        f"({fc['30']['Calibrated twin (physiology only)']:.0f} vs {fc['30']['Persistence (no change)']:.0f} mg/dL at "
        f"+30 min): its dynamics are too fast for these closely managed patients. The hybrid ML forecaster was best "
        f"({fc['30']['GlucoTwin hybrid (ML, out-of-fold)']:.0f} mg/dL).",
    ]


def update_readme(m: dict) -> None:
    path = ROOT / "README.md"
    if not path.exists():
        return
    s = path.read_text(encoding="utf-8")
    for tag, body in (("RESULTS", results_markdown(m)), ("TAKEAWAYS", takeaways_markdown(m)),
                      ("GLANCE", glance_markdown(m))):
        block = f"<!-- {tag}:START -->\n{body}\n<!-- {tag}:END -->"
        s = re.sub(rf"<!-- {tag}:START -->.*?<!-- {tag}:END -->", lambda _, b=block: b, s, flags=re.S)
    path.write_text(s, encoding="utf-8")


# --------------------------------------------------------------------------------------
# Presentation
# --------------------------------------------------------------------------------------
def presentation(m: dict) -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.util import Inches, Pt

    def rgb(h):
        return RGBColor.from_string(h.lstrip("#"))

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]
    W = prs.slide_width

    def text(slide, x, y, w, h, s, size=16, color=INK, bold=False):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, line in enumerate(s if isinstance(s, list) else [s]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.font.size = Pt(size)
            p.font.bold = bold or (isinstance(s, list) and i == 0 and line in ("Limitations", "Next steps", "Stack",
                                                                               "Our question"))
            p.font.color.rgb = rgb(color)
            p.font.name = "Segoe UI"
            p.space_after = Pt(6)
        return tb

    def fill(shape, color):
        shape.fill.solid()
        shape.fill.fore_color.rgb = rgb(color)
        shape.line.fill.background()

    def slide(title, kicker=None):
        s = prs.slides.add_slide(blank)
        fill(s.shapes.add_shape(1, 0, 0, W, Inches(0.12)), BLUE)
        if kicker:
            text(s, 0.6, 0.35, 12, 0.4, kicker.upper(), 12, BLUE, True)
        text(s, 0.6, 0.65, 12.2, 0.9, title, 28, INK, True)
        text(s, 0.6, 7.0, 12, 0.35, f"GlucoTwin · {TEAM} · Digital Twin Challenge 2026", 10, MUTED)
        return s

    def image(s, name, x, y, w, h=None):
        """Place an image inside the box (x, y, w, h), keeping its aspect ratio, centred horizontally.
        The box defaults to everything above the footer."""
        from PIL import Image

        p = FIGS / name if (FIGS / name).exists() else DOCS / name
        if not p.exists():
            return
        h = h if h is not None else 6.85 - y
        iw, ih = Image.open(p).size
        scale = min(w / iw, h / ih)
        fw, fh = iw * scale, ih * scale
        s.shapes.add_picture(str(p), Inches(x + (w - fw) / 2), Inches(y), Inches(fw), Inches(fh))

    def stat(s, x, y, value, label, color=BLUE):
        text(s, x, y, 2.9, 0.8, value, 34, color, True)
        text(s, x, y + 0.8, 2.9, 0.8, label, 12, INK2)

    ab = m["ablations"]
    fin = ab["GlucoTwin: fusion + digital twin"]
    cgm = ab["CGM only"]
    wear_ = ab["Wearables only (dynamic)"]
    ci_cgm = m["auroc_ci"]["CGM only"]
    fc = m["forecast"]

    # 1 Title
    s = prs.slides.add_slide(blank)
    fill(s.shapes.add_shape(1, 0, 0, W, prs.slide_height), "#0d366b")
    text(s, 0.8, 1.6, 11.5, 1.2, "GlucoTwin", 54, "#ffffff", True)
    text(s, 0.8, 2.7, 11.5, 1.4, "A hybrid digital twin that warns of glucose spikes and lows in Type 2 Diabetes "
         "before they happen, by fusing EHR, CGM and wearable data, validated on real patients", 22, "#cde2fb")
    text(s, 0.8, 4.6, 11.5, 1.2, [f"{TEAM}  ·  {COLLEGE}", "Digital Twin Challenge 2026 · Happiest Health · "
                                  "Reimagining & Reforming Healthcare in India Summit"], 14, "#ffffff")

    # 2 Problem
    s = slide("Diabetes in India is managed by looking backwards", "The problem")
    text(s, 0.6, 1.7, 6.3, 5, [
        "• India has ~101 million adults with diabetes and ~136 million with prediabetes (ICMR-INDIAB, 2023).",
        "• High-carbohydrate meals (rice, roti, dosa, biryani) drive sharp post-meal glucose excursions, "
        "linked to vascular damage even when HbA1c looks acceptable.",
        "• HbA1c is measured every ~3 months: it averages away the daily spikes.",
        "• CGMs show what already happened, not what is about to happen.",
        "• Doctors see a patient for minutes every few months, with no way to test "
        "'what if this patient walked after dinner?'"], 16, INK2)
    fill(s.shapes.add_shape(1, Inches(7.4), Inches(1.9), Inches(5.3), Inches(3.6)), "#f3f2ee")
    text(s, 7.7, 2.1, 4.8, 3.3, ["Our question", "Can a personalised virtual replica of a patient, continuously updated "
                                 "from wearables and anchored in their health record, warn of a glucose spike early "
                                 "enough to act, and show which action would help?"], 18, INK)

    # 3 Architecture
    s = slide("A three-layer hybrid digital twin", "Our solution")
    image(s, "architecture.png", 0.6, 1.5, w=12.1)

    # 4 Data
    s = slide("Sandbox-compliant data: a causally consistent synthetic cohort", "Data")
    text(s, 0.6, 1.6, 5.4, 5, [
        "• 120 synthetic Indian T2D patients × 28 days at 5-minute resolution (≈1M rows).",
        "• Static EHR: demographics, labs, comorbidities, 6 drug classes, TCF7L2/KCNQ1 risk alleles, polygenic score.",
        "• Dynamic streams: CGM, steps, heart rate, HRV, sleep stages, partial meal logs.",
        "• Generated causally: genetics + meds → physiology → glucose → sensors and labs, so HbA1c agrees with CGM.",
        "• Hidden drivers the model must infer: unlogged meals (≈1 in 3), stress (HR↑ HRV↓), sleep debt.",
        "• Realistic noise: CGM lag, ~7% MARD, dropouts, sensor change, HRV gaps.",
        "• Real-data path built and tested: Synthea EHR importer, Apple Health importer, FHIR R4 export."], 14, INK2)
    image(s, "cohort_validation.png", 6.2, 1.8, w=6.7)

    # 5 Mechanistic twin
    s = slide("Layer ①: a personalised physiological twin", "Mechanistic model")
    text(s, 0.6, 1.6, 12, 1.6, [
        "dG/dt = −(SG + X·m)·G + SG·Gb + Ra(meals) + dawn − exercise − SGLT2 excretion   ·   "
        "dX/dt = −p2·X + p2·SI·(I + I_exo)   ·   dI/dt = −n·I + β·(G − Gb)⁺ + s0",
        "EHR sets the priors (weight, FPG, BMI, duration, medications). One onboarding week of CGM + meals + steps "
        "personalises insulin sensitivity, β-cell response, basal glucose and dawn effect (MAP fit). "
        "CGM is assimilated every 5 minutes."], 13, INK2)
    image(s, "twin_recovery.png", 0.9, 3.25, w=11.5)

    # 6 Results headline
    s = slide("Results on 24 patients the model never saw", "Performance")
    stat(s, 0.6, 1.7, f"{fin['auroc']:.3f}", "AUROC for predicting a spike\nwithin the next 2 hours")
    stat(s, 3.7, 1.7, f"{fin['event_sensitivity']:.0%}", "of spikes warned about\nin advance")
    stat(s, 6.8, 1.7, f"{fin['median_lead_min']:.0f} min", "median warning time\nbefore the spike starts", ORANGE)
    stat(s, 9.9, 1.7, f"{fin['false_alerts_per_day']:.2f}", "false alerts per\npatient-day", AQUA)
    image(s, "ablation.png", 0.6, 3.6, w=12.1)

    # 7 Fusion story
    s = slide("Why fuse? Each data stream earns its place", "Ablation")
    text(s, 0.6, 1.6, 12, 1.6, [
        f"• CGM alone: AUROC {cgm['auroc']:.3f}, {cgm['false_alerts_per_day']:.2f} false alerts/day. Adding meals, "
        f"activity, sleep, HR/HRV and personal routine: {wear_['auroc']:.3f} and {wear_['false_alerts_per_day']:.2f}/day "
        f"(AUROC gain {ci_cgm['diff']:+.3f}, 95% CI {ci_cgm['lo']:+.3f} to {ci_cgm['hi']:+.3f}).",
        "• Once wearable history exists, static EHR and twin features add no measurable discrimination "
        "(differences within ±0.002). The twin's role is simulation and explanation, not raw accuracy."],
         14, INK2)
    image(s, "by_history.png", 0.6, 3.0, w=12.1)

    # 7b Real patients
    rd = real_metrics()
    if rd:
        s = slide(f"Validated on {rd['cohort']['subjects']} real people with type 2 diabetes", "Real-world evidence")
        g = rd["diff_ci"]["GlucoTwin (+ digital twin) vs CGM only"]
        text(s, 0.6, 1.55, 12.2, 1.4, [
            "ShanghaiT2DM (Sci Data 2023, CC BY 4.0): CGM + clinical record + meals + insulin doses. "
            "5-fold cross-validation grouped by person.",
            f"Fusion beats CGM alone ({g['diff']:+.3f} AUROC, 95% CI {g['lo']:+.3f} to {g['hi']:+.3f}). Trained only on "
            f"synthetic data: {rd['zero_shot']['auroc']:.3f}. Honest negatives: static EHR slightly hurt; the "
            "physics-only twin forecast is worse than 'no change' on these patients."], 13, INK2)
        image(s, "real_data.png", 0.6, 3.0, w=12.1)

    # 7c Lows
    hy = hypo_metrics()
    if hy:
        s = slide("Both directions of danger: predicting lows on real patients", "Hypoglycaemia")
        g = hy["auroc"]["GlucoTwin (+ digital twin)"]
        rule = hy["auroc"]["Rule: 'urgent low soon' (30-min extrapolation)"]
        text(s, 0.6, 1.55, 12.2, 1.4, [
            f"Lows (<70 mg/dL) are the dangerous side of treatment, especially with insulin or sulfonylureas. Trained and "
            f"validated on {hy['cohort']['people']} real people ({hy['cohort']['low_episodes']} low episodes), "
            "with person-grouped cross-validation.",
            f"GlucoTwin AUROC {g['auroc']:.3f} vs {rule['auroc']:.3f} for the 'urgent low soon' rule that CGMs use. "
            "Mixed but useful: GlucoTwin catches more lows when alarms must be rare (alarm fatigue), the simple rule "
            "wins when alarms can be frequent, so we use both: the rule as a safety net, GlucoTwin for earlier, quieter "
            "warnings."], 13, INK2)
        image(s, "hypo_real_data.png", 0.6, 3.0, w=12.1)

    # 8 Early warning
    s = slide("Early, actionable warnings", "Lead time")
    image(s, "example_patient.png", 0.5, 1.6, w=7.6)
    image(s, "lead_time.png", 8.3, 1.9, w=4.6)

    # 9 Forecast & calibration
    s = slide("Trustworthy numbers: calibrated risk and accurate forecasts", "Forecasting")
    text(s, 0.6, 1.55, 12, 1, f"Hybrid forecast error at +30 / +60 / +120 min: "
         f"{fc['30']['GlucoTwin hybrid']['rmse']:.0f} / {fc['60']['GlucoTwin hybrid']['rmse']:.0f} / "
         f"{fc['120']['GlucoTwin hybrid']['rmse']:.0f} mg/dL, roughly half the error of the physiology-only twin "
         "or a 'no change' guess.", 14, INK2)
    image(s, "calibration_forecast.png", 0.9, 2.4, w=11.5)

    # 10 Explainability
    s = slide("Explainable by design", "Why did it alert?")
    image(s, "group_importance.png", 0.6, 1.6, w=6.2)
    text(s, 7.2, 1.7, 5.6, 5, [
        "Every alert comes with its top drivers, for example:",
        "▲ Personal routine: spiked above 180 around this time on 3 of the last 3 days",
        "▲ Meals: 75 g carbs logged 20 min ago",
        "▲ Digital twin: glucose 28 mg/dL above twin expectation, possible unlogged meal",
        "▲ Sleep: slept 4.9 h last night (−1.8 h vs usual)",
        "",
        "…plus a twin-tested suggestion, e.g. 'a 15-minute brisk walk now lowers the predicted peak by ~20 mg/dL'."],
         14, INK2)

    # 11 Dashboard
    s = slide("How a doctor uses the twin", "Clinician dashboard")
    shots = [p for p in ("dashboard_ward.png", "dashboard_patient.png") if (DOCS / p).exists()]
    for i, p in enumerate(shots):
        image(s, p, 0.6 + i * 6.2, 1.55, w=6.0, h=4.0)
    text(s, 0.6, 5.75, 12, 1.2, "Ward overview (risk-sorted)  →  Virtual patient (CGM, forecast, risk, activity, HR/HRV, "
         "sleep, EHR, AGP, twin phenotype)  →  What-if simulator (Indian meal presets, carb quality, post-meal walk, "
         "sleep)  →  Model performance", 14, INK2)

    # 11a Twin overview (body map)
    if (DOCS / "dashboard_body.png").exists():
        s = slide("The twin at a glance: the patient's body, live", "Twin overview")
        image(s, "dashboard_body.png", 0.6, 1.4, w=12.1, h=5.5)

    # 11b Patient phone
    if (DOCS / "dashboard_phone.png").exists():
        s = slide("What the patient sees: one number, one reason, one action", "Patient app")
        image(s, "dashboard_phone.png", 1.2, 1.4, w=6.0, h=5.4)
        text(s, 8.4, 1.7, 4.4, 5, [
            "The clinician gets the full console; the patient gets a phone screen:",
            "• glucose now + trend",
            "• the prediction in plain words (English or Hindi)",
            "• the top reason, e.g. 'You had a meal recently'",
            "• one action the twin has tested: a 15-minute walk",
            "• lows take priority over highs",
            "",
            "The 'I'll do it / Remind me' taps would feed adherence back to the care team."], 14, INK2)

    # 12 What-if
    s = slide("What-if: personalised treatment exploration", "Simulation")
    image(s, "dashboard_whatif.png", 0.6, 1.6, w=7.4)
    text(s, 8.3, 1.7, 4.5, 5, [
        "The calibrated twin is run forward from the patient's current state:",
        "• Which meal? (rice + sambar vs ragi mudde)",
        "• Walk after dinner? How long, how soon?",
        "• Effect of a short night's sleep",
        "",
        "For counselling and shared decision-making, not dosing."], 15, INK2)

    # 13 Engineering
    s = slide("Built like a product, not a notebook", "Engineering")
    text(s, 0.6, 1.6, 6, 5, [
        "• Real-time engine + FastAPI: POST observations → GET assessment / what-if",
        "• Replay script streams a patient like a device gateway",
        "• Identical feature code online and offline (tested to match)",
        "• Leakage test: changing the future never changes today's features",
        "• Patient-level train/val/test split; thresholds chosen on validation patients",
        "• Alerts evaluated as real notifications (re-sent hourly while risk stays high)",
        "• One command rebuilds everything: python scripts/build_all.py",
        "• New-patient lifecycle over the API: register → stream → calibrate → FHIR RiskAssessment",
        "• docker compose up runs dashboard + API"], 15, INK2)
    text(s, 7.0, 1.6, 5.8, 5, [
        "Stack",
        "Python · NumPy · pandas · SciPy · Numba (ODE kernels) · scikit-learn (HistGradientBoosting, isotonic "
        "calibration) · FastAPI · Streamlit · Plotly · pytest", "", "License: MIT"], 15, INK2)

    # 14 Limitations & next steps
    s = slide("Limitations and next steps", "Honest assessment")
    text(s, 0.6, 1.6, 6, 5, [
        "Limitations",
        "• Main results are synthetic; the real-data check (96 Chinese T2D patients, CGM but no other "
        "wearables) confirms the fusion gain but not the twin's forecasting.",
        "• What-if simulations are not validated against real interventions.",
        "• The low-glucose model is trained only on the real Shanghai cohort (lows are too rare in the simulation).",
        "• The dawn-effect parameter is not identifiable from one week.",
        "• Twin features add no measurable accuracy over the fused ML model, even when trained on only "
        "6 patients; the twin's value is what-if simulation and interpretability.",
        "• Once wearable history exists, static EHR adds little discrimination.",
        ], 15, INK2)
    text(s, 7.0, 1.6, 5.8, 5, [
        "Next steps",
        "• Prospective Indian cohort with CGM + smartwatch (steps, HR, HRV, sleep) to test the full fusion.",
        "• Connect the FHIR R4 export to an ABDM health-information gateway (export itself is built).",
        "• Validate the low-glucose model on Indian patients and add nocturnal-low alerts.",
        "• Prospective pilot with clinicians, then the CDSCO SaMD pathway."], 15, INK2)

    # 15 Thanks
    s = prs.slides.add_slide(blank)
    fill(s.shapes.add_shape(1, 0, 0, W, prs.slide_height), "#0d366b")
    text(s, 0.8, 2.6, 11.5, 1.2, "Thank you", 48, "#ffffff", True)
    text(s, 0.8, 3.7, 11.5, 1.2, f"{TEAM} · {COLLEGE} · github.com/maxlovesher/twin_KIIT", 18, "#cde2fb")
    text(s, 0.8, 4.4, 11.5, 1.0, "Live demo: glucotwin-kiit.streamlit.app", 18, "#ffffff", True)

    out = DOCS / "GlucoTwin_presentation.pptx"
    prs.save(out)
    return out


def pptx_to_pdf(pptx: Path) -> None:
    """Export via PowerPoint (Windows) or LibreOffice if available; skipped otherwise."""
    pdf = pptx.with_suffix(".pdf")
    ps = (f"$p = New-Object -ComObject PowerPoint.Application; "
          f"$d = $p.Presentations.Open('{pptx}', $true, $false, $false); "
          f"$d.SaveAs('{pdf}', 32); $d.Close(); $p.Quit()")
    try:
        subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True, timeout=180, capture_output=True)
        return
    except Exception:
        pass
    try:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(pptx.parent), str(pptx)],
                       check=True, timeout=180, capture_output=True)
    except Exception:
        print("PDF export skipped (needs PowerPoint or LibreOffice)")


def main() -> None:
    m = metrics()
    architecture()
    update_readme(m)
    out = presentation(m)
    pptx_to_pdf(out)
    print("wrote", out, "and docs/architecture.pdf")


if __name__ == "__main__":
    main()
